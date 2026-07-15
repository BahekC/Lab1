from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

FONT_NAME = "Times New Roman"
BASE_SIZE = Pt(14)


def set_run_font(run, size=BASE_SIZE, bold=None, italic=None):
    run.font.name = FONT_NAME
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    run.font.size = size
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=80, bottom=80, end=80):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_run_font(run, Pt(10))
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def configure_document(doc: DocumentType, footer_text: str):
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_NAME
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    normal.font.size = BASE_SIZE
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)
    pf.first_line_indent = Cm(1.25)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for style_name, size in (("Title", Pt(18)), ("Heading 1", Pt(16)), ("Heading 2", Pt(14)), ("Heading 3", Pt(14))):
        st = styles[style_name]
        st.font.name = FONT_NAME
        st._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
        st.font.size = size
        st.font.bold = True
        st.paragraph_format.space_before = Pt(10)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.first_line_indent = Cm(0)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hr = hp.add_run(footer_text)
    set_run_font(hr, Pt(9))

    footer = section.footer
    fp = footer.paragraphs[0]
    add_page_number(fp)


def add_text(doc: DocumentType, text: str, *, bold=False, italic=False, align=None, indent=True, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.first_line_indent = Cm(1.25) if indent else Cm(0)
    p.paragraph_format.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.keep_together = keep
    r = p.add_run(text)
    set_run_font(r, BASE_SIZE, bold=bold, italic=italic)
    return p


def add_bullets(doc: DocumentType, items: Iterable[str], ordered=False):
    for idx, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.5)
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(0)
        prefix = f"{idx}. " if ordered else "• "
        r = p.add_run(prefix + item)
        set_run_font(r)


def add_numbered(doc: DocumentType, items: Iterable[str]):
    add_bullets(doc, items, ordered=True)


def add_heading(doc: DocumentType, text: str, level=1):
    p = doc.add_heading(text, level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    for r in p.runs:
        set_run_font(r, Pt(16) if level == 1 else Pt(14), bold=True)
    return p


def add_table(doc: DocumentType, headers: Sequence[str], rows: Sequence[Sequence[str]], widths: Sequence[float] | None = None, caption: str | None = None):
    if caption:
        add_text(doc, caption, bold=True, indent=False, keep=True)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, "D9E2F3")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(str(h))
        set_run_font(r, Pt(11), bold=True)
    for row_data in rows:
        row = table.add_row()
        for i, value in enumerate(row_data):
            cell = row.cells[i]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            set_run_font(r, Pt(11))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_title_page(doc: DocumentType, lecture_no: int, lecture_total: int, title: str, module_hours: int = 10):
    add_heading(doc, "1. Титульный лист.", 1)
    lines = [
        ("ОБЩЕСТВО С ОГРАНИЧЕННОЙ ОТВЕТСТВЕННОСТЬЮ", False, Pt(14)),
        ("«ЭКСПАЛЬЯНС»", True, Pt(16)),
        ("", False, Pt(14)),
        ("ДОПОЛНИТЕЛЬНАЯ ПРОФЕССИОНАЛЬНАЯ ПРОГРАММА", True, Pt(14)),
        ("ПОВЫШЕНИЯ КВАЛИФИКАЦИИ", True, Pt(14)),
        ("«ВЕДЕНИЕ БИЗНЕСА С КИТАЕМ»", True, Pt(17)),
        ("", False, Pt(14)),
        ("Модуль 5. Электронные торговые площадки Китая", True, Pt(15)),
        (f"Лекция {lecture_no} из {lecture_total}", True, Pt(14)),
        ("", False, Pt(14)),
        (title.upper(), True, Pt(18)),
        ("", False, Pt(14)),
        ("Объем программы: 144 академических часа", False, Pt(14)),
        (f"Трудоемкость модуля: {module_hours} академических часов", False, Pt(14)),
        ("Категория слушателей: руководители организаций, предприниматели, специалисты ВЭД", False, Pt(14)),
        ("Форма обучения: с применением электронного обучения и дистанционных образовательных технологий", False, Pt(14)),
        ("", False, Pt(14)),
        ("Учебно-методический материал для дистанционного обучения взрослых слушателей ДПО", False, Pt(14)),
        ("", False, Pt(14)),
        ("г. Камбарка", False, Pt(14)),
        ("2026", False, Pt(14)),
    ]
    for text, bold, size in lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(text if text else " ")
        set_run_font(r, size, bold=bold)
    doc.add_page_break()


def add_section(doc: DocumentType, number: int, title: str, paragraphs: Sequence[str] | None = None, bullets: Sequence[str] | None = None, page_break=False):
    if page_break:
        doc.add_page_break()
    add_heading(doc, f"{number}. {title}.", 1)
    if paragraphs:
        for t in paragraphs:
            add_text(doc, t)
    if bullets:
        add_bullets(doc, bullets)


def add_main_subsection(doc: DocumentType, index: str, title: str, paragraphs: Sequence[str], bullets: Sequence[str] | None = None, table=None, mini_case=None):
    add_heading(doc, f"7.{index}. {title}", 2)
    for t in paragraphs:
        add_text(doc, t)
    if bullets:
        add_bullets(doc, bullets)
    if table:
        add_table(doc, table[0], table[1], caption=table[2] if len(table) > 2 else None)
    if mini_case:
        add_text(doc, "Мини-кейс", bold=True, indent=False, keep=True)
        for t in mini_case:
            add_text(doc, t)


def add_practical_subsection(doc: DocumentType, index: str, title: str, paragraphs: Sequence[str], bullets: Sequence[str] | None = None, table=None):
    add_heading(doc, f"8.{index}. {title}", 2)
    for t in paragraphs:
        add_text(doc, t)
    if bullets:
        add_bullets(doc, bullets)
    if table:
        add_table(doc, table[0], table[1], caption=table[2] if len(table) > 2 else None)


def add_learning_results(doc: DocumentType, know, be_able, possess):
    add_heading(doc, "5. Планируемые результаты обучения.", 1)
    add_heading(doc, "Знать", 2)
    add_bullets(doc, know)
    add_heading(doc, "Уметь", 2)
    add_bullets(doc, be_able)
    add_heading(doc, "Владеть", 2)
    add_bullets(doc, possess)


def add_questions(doc: DocumentType, questions: Sequence[str]):
    add_heading(doc, "9. Контрольные вопросы.", 1)
    add_numbered(doc, questions)


def add_assignment(doc: DocumentType, paragraphs: Sequence[str], tasks: Sequence[str]):
    add_heading(doc, "10. Практическое задание.", 1)
    for t in paragraphs:
        add_text(doc, t)
    add_numbered(doc, tasks)


def add_criteria(doc: DocumentType, rows):
    add_heading(doc, "11. Критерии оценивания.", 1)
    add_text(doc, "Работа оценивается по совокупности критериев. Зачет рекомендуется выставлять при наборе не менее 70 баллов из 100 при условии отсутствия критической ошибки, способной привести к платежу ненадлежащему лицу, заказу неподтвержденного товара либо утрате доказательств договоренности.")
    add_table(doc, ("Критерий", "Максимум", "Признаки качественного выполнения"), rows, caption="Таблица критериев оценивания")


def add_terms(doc: DocumentType, terms):
    add_heading(doc, "12. Термины и определения.", 1)
    add_table(doc, ("Термин", "Определение"), terms, caption="Ключевые термины лекции")


def add_recommendations(doc: DocumentType, paragraphs: Sequence[str], bullets: Sequence[str]):
    add_heading(doc, "13. Методические рекомендации слушателю.", 1)
    for t in paragraphs:
        add_text(doc, t)
    add_bullets(doc, bullets)


def add_conclusion(doc: DocumentType, paragraphs: Sequence[str]):
    add_heading(doc, "14. Заключение.", 1)
    for t in paragraphs:
        add_text(doc, t)


def add_sources(doc: DocumentType, sources: Sequence[str]):
    add_heading(doc, "15. Список источников.", 1)
    add_numbered(doc, sources)


def validate_and_save(doc: DocumentType, output_path: Path, expected_title: str):
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    check = Document(output_path)
    all_text = "\n".join(p.text for p in check.paragraphs)
    required = [
        "1. Титульный лист.", "2. Аннотация.", "3. Цель лекции.", "4. Задачи лекции.",
        "5. Планируемые результаты обучения.", "Знать", "Уметь", "Владеть", "6. План лекции.",
        "7. Основной текст лекции.", "8. Практико-ориентированный блок.", "9. Контрольные вопросы.",
        "10. Практическое задание.", "11. Критерии оценивания.", "12. Термины и определения.",
        "13. Методические рекомендации слушателю.", "14. Заключение.", "15. Список источников.",
    ]
    missing = [x for x in required if x not in all_text]
    if missing:
        raise RuntimeError(f"Missing headings in {output_path.name}: {missing}")
    if expected_title.upper() not in all_text.upper():
        raise RuntimeError(f"Title missing in {output_path.name}")
    words = len(all_text.split())
    if words < 6000:
        raise RuntimeError(f"Document is too short: {words} words in {output_path.name}")
    return words
